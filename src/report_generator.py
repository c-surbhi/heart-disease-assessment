from docx import Document
from docx.shared import Inches

import pandas as pd
import os


REPORT_PATH = (
    "outputs/reports/"
    "model_evaluation_report.docx"
)


def add_image(
    document,
    image_path,
    width=5
):

    if os.path.exists(image_path):

        document.add_picture(
            image_path,
            width=Inches(width)
        )


def add_dataframe(
    document,
    dataframe,
    title
):

    document.add_heading(
        title,
        level=2
    )

    rows, cols = dataframe.shape

    table = document.add_table(
        rows=rows + 1,
        cols=cols
    )

    table.style = "Table Grid"

    for col_num, column in enumerate(
        dataframe.columns
    ):
        table.cell(
            0,
            col_num
        ).text = str(column)

    for row in range(rows):

        for col in range(cols):

            table.cell(
                row + 1,
                col
            ).text = str(
                dataframe.iloc[row, col]
            )


def add_metrics_section(document):

    logistic = pd.read_csv(
        "outputs/reports/"
        "logistic_regression_metrics.csv"
    )

    rf = pd.read_csv(
        "outputs/reports/"
        "random_forest_metrics.csv"
    )

    document.add_heading(
        "Model Evaluation Metrics",
        level=1
    )

    add_dataframe(
        document,
        logistic,
        "Logistic Regression Metrics"
    )

    add_dataframe(
        document,
        rf,
        "Random Forest Metrics"
    )


def add_decile_section(document):

    logistic = pd.read_csv(
        "outputs/reports/"
        "logistic_deciles.csv"
    )

    rf = pd.read_csv(
        "outputs/reports/"
        "rf_deciles.csv"
    )

    document.add_heading(
        "Decile Analysis",
        level=1
    )

    add_dataframe(
        document,
        logistic,
        "Logistic Regression Deciles"
    )

    add_dataframe(
        document,
        rf,
        "Random Forest Deciles"
    )


def add_threshold_section(document):

    summary = pd.read_csv(
        "outputs/reports/"
        "profit_summary.csv"
    )

    document.add_heading(
        "Threshold Optimization",
        level=1
    )

    add_dataframe(
        document,
        summary,
        "Optimal Threshold Summary"
    )

def add_visualizations(document):

    document.add_heading(
        "Visualizations",
        level=1
    )

    add_image(
        document,
        "outputs/plots/cumulative_gain.png"
    )

    add_image(
        document,
        "outputs/plots/lift_chart.png"
    )

    add_image(
        document,
        "outputs/plots/profit_curve_logistic.png"
    )

    add_image(
        document,
        "outputs/plots/profit_curve_random_forest.png"
    )

    add_image(
        document,
        "outputs/plots/logistic_regression_roc.png"
    )

    add_image(
        document,
        "outputs/plots/random_forest_roc.png"
    )

def add_business_summary(document):

    document.add_heading(
        "Business Recommendation",
        level=1
    )

    summary = pd.read_csv(
        "outputs/reports/"
        "profit_summary.csv"
    )

    best_model = (
        summary
        .sort_values(
            "Profit",
            ascending=False
        )
        .iloc[0]
    )

    recommendation = f"""
Recommended Model:
{best_model['Model']}

Reason:
This model generated the
highest business profit after
threshold optimization.

The recommendation is based
on business value rather than
accuracy alone.

The selected threshold
maximizes:

- True Positives
- Cost Savings
- Event Capture

while minimizing expensive
False Negatives.
"""

    document.add_paragraph(
        recommendation
    )


def main():

    document = Document()

    document.add_heading(
        "Heart Disease Prediction",
        level=0
    )

    document.add_paragraph(
        """
This report compares
Logistic Regression and
Random Forest models for
Heart Disease Prediction.

Analysis includes:

• Model Evaluation

• Decile Analysis

• Lift Analysis

• Threshold Optimization

• Business Profit Analysis
"""
    )

    add_metrics_section(
        document
    )

    add_decile_section(
        document
    )

    add_threshold_section(
        document
    )

    add_visualizations(
        document
    )

    add_business_summary(
        document
    )

    document.save(
        REPORT_PATH
    )

    print(
        f"Report Saved: "
        f"{REPORT_PATH}"
    )


if __name__ == "__main__":
    main()
